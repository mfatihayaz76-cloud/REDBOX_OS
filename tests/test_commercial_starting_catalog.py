import hashlib
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from database.recipe_import_engine import (
    ALL_HEADERS,
    katalog_dosyasi_on_kontrol,
    katalog_satirlarini_oku,
)
from tools.export_commercial_starting_catalog import (
    export_catalog,
)


class CommercialStartingCatalogTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.database_path = (
            PROJECT_ROOT
            / "database"
            / "redbox_os.db"
        )

    def database_sha(self):
        return hashlib.sha256(
            self.database_path.read_bytes()
        ).hexdigest()

    def create_empty_catalog_sandbox(self):
        live = sqlite3.connect(
            f"file:{self.database_path}?mode=ro",
            uri=True,
        )
        live.row_factory = sqlite3.Row

        sandbox = sqlite3.connect(":memory:")
        sandbox.row_factory = sqlite3.Row

        schema_rows = live.execute(
            """
            SELECT type, name, sql
            FROM sqlite_master
            WHERE sql IS NOT NULL
              AND name NOT LIKE 'sqlite_%'
            ORDER BY
                CASE type
                    WHEN 'table' THEN 1
                    WHEN 'index' THEN 2
                    WHEN 'trigger' THEN 3
                    ELSE 4
                END,
                name
            """
        ).fetchall()

        for row in schema_rows:
            sandbox.execute(row["sql"])

        columns = [
            row["name"]
            for row in live.execute(
                "PRAGMA table_info(hammaddeler)"
            ).fetchall()
        ]
        materials = live.execute(
            "SELECT * FROM hammaddeler"
        ).fetchall()

        column_sql = ", ".join(
            f'"{column}"'
            for column in columns
        )
        placeholders = ", ".join(
            "?"
            for _ in columns
        )

        sandbox.executemany(
            (
                f"INSERT INTO hammaddeler ({column_sql}) "
                f"VALUES ({placeholders})"
            ),
            [
                tuple(row[column] for column in columns)
                for row in materials
            ],
        )
        sandbox.commit()
        live.close()

        return sandbox

    def test_real_catalog_export_and_format_parity(self):
        before_sha = self.database_sha()

        with tempfile.TemporaryDirectory() as temp_dir:
            result = export_catalog(
                self.database_path,
                Path(temp_dir),
            )

            csv_headers, csv_rows = (
                katalog_satirlarini_oku(
                    result["csv_path"]
                )
            )
            xlsx_headers, xlsx_rows = (
                katalog_satirlarini_oku(
                    result["xlsx_path"]
                )
            )

        self.assertEqual(
            before_sha,
            self.database_sha(),
        )
        self.assertEqual(
            csv_headers,
            list(ALL_HEADERS),
        )
        self.assertEqual(
            xlsx_headers,
            list(ALL_HEADERS),
        )
        self.assertEqual(result["urun_sayisi"], 1)
        self.assertEqual(result["recete_sayisi"], 2)
        self.assertEqual(result["satir_sayisi"], 16)
        self.assertEqual(len(csv_rows), 16)
        self.assertEqual(len(xlsx_rows), 16)

        def normalize(rows):
            return [
                {
                    header: (
                        ""
                        if row.get(header) is None
                        else str(row.get(header))
                    )
                    for header in ALL_HEADERS
                }
                for row in rows
            ]

        self.assertEqual(
            normalize(csv_rows),
            normalize(xlsx_rows),
        )

        recipe_contract = {
            (
                row["urun_kodu"],
                row["recete_kodu"],
                row["revizyon_no"],
                row["durum"],
            )
            for row in csv_rows
        }

        self.assertEqual(
            recipe_contract,
            {
                (
                    "LP001",
                    "LP001-REC",
                    "00",
                    "ARSIV",
                ),
                (
                    "LP001",
                    "LP001-REC",
                    "01",
                    "AKTIF",
                ),
            },
        )

        active_rows = [
            row
            for row in csv_rows
            if row["revizyon_no"] == "01"
        ]
        active_formula = {
            row["hammadde_adi"]: float(
                row["miktar_kg"]
            )
            for row in active_rows
        }

        self.assertEqual(
            active_formula,
            {
                "Patates Unu": 5.2,
                "Nişasta": 2.56,
                "Mısır Unu": 1.28,
                (
                    "Metilselüloz Benecel "
                    "A4M E461"
                ): 0.12,
                "Tavuk Çeşnisi": 0.32,
                "Sarımsak Tozu": 0.064,
                "Karabiber": 0.024,
                "Tuz": 0.144,
            },
        )
        self.assertEqual(
            {
                float(row["parti_teorik_kg"])
                for row in active_rows
            },
            {20.412},
        )
        self.assertEqual(
            {
                float(row["proses_suyu_kg"])
                for row in active_rows
            },
            {10.7},
        )
        self.assertEqual(
            {
                row["gecerlilik_tarihi"]
                for row in active_rows
            },
            {"20.07.2027"},
        )

    def test_xlsx_professional_layout_contract(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            result = export_catalog(
                self.database_path,
                Path(temp_dir),
            )
            workbook = load_workbook(
                result["xlsx_path"],
                read_only=False,
                data_only=False,
            )

            try:
                sheet = workbook["RECETE_KATALOGU"]

                self.assertEqual(
                    sheet.freeze_panes,
                    "A2",
                )
                self.assertEqual(
                    sheet.auto_filter.ref,
                    "A1:R5000",
                )
                self.assertEqual(
                    sheet.row_dimensions[1].height,
                    48,
                )
                self.assertEqual(
                    sheet.row_dimensions[2].height,
                    32,
                )
                self.assertGreaterEqual(
                    sheet.column_dimensions["D"].width,
                    42,
                )
                self.assertGreaterEqual(
                    sheet.column_dimensions["J"].width,
                    38,
                )
                self.assertGreaterEqual(
                    sheet.column_dimensions["Q"].width,
                    48,
                )
                self.assertGreaterEqual(
                    sheet.column_dimensions["R"].width,
                    54,
                )
                self.assertTrue(
                    sheet["D2"].alignment.wrap_text
                )
                self.assertTrue(
                    sheet["J2"].alignment.wrap_text
                )
                self.assertEqual(
                    sheet["H2"].number_format,
                    "0.000",
                )
                self.assertEqual(
                    sheet["E2"].number_format,
                    "@",
                )
            finally:
                workbook.close()

    def test_csv_and_xlsx_pass_empty_sandbox_dry_run(self):
        before_sha = self.database_sha()
        sandbox = self.create_empty_catalog_sandbox()

        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                result = export_catalog(
                    self.database_path,
                    Path(temp_dir),
                )

                csv_report = katalog_dosyasi_on_kontrol(
                    sandbox,
                    result["csv_path"],
                )
                xlsx_report = katalog_dosyasi_on_kontrol(
                    sandbox,
                    result["xlsx_path"],
                )

            self.assertTrue(csv_report["gecerli"])
            self.assertTrue(xlsx_report["gecerli"])
            self.assertEqual(
                csv_report["ozet"],
                xlsx_report["ozet"],
            )
            self.assertEqual(
                csv_report["ozet"],
                {
                    "satir_sayisi": 16,
                    "urun_sayisi": 1,
                    "recete_sayisi": 2,
                    "kalem_sayisi": 16,
                    "hata_sayisi": 0,
                    "uyari_sayisi": 0,
                },
            )
        finally:
            sandbox.close()

        self.assertEqual(
            before_sha,
            self.database_sha(),
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
